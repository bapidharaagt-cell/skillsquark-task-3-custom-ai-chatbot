"""
SkillsQuark Task 3 – Custom AI Chatbot

A beginner-friendly document question-answering chatbot using:
- Streamlit
- LangChain
- OpenAI embeddings + chat model
- FAISS vector store
- PyPDF for PDF extraction

The chatbot accepts PDF, TXT, CSV, and DOCX files, indexes their content,
and answers questions using retrieved document context.

Run:
    streamlit run app.py
"""

import os
import tempfile
from pathlib import Path

import pandas as pd
import streamlit as st
from dotenv import load_dotenv
from pypdf import PdfReader

from langchain_core.documents import Document
from langchain_openai import ChatOpenAI, OpenAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain_core.prompts import ChatPromptTemplate


load_dotenv()

st.set_page_config(
    page_title="Custom AI Document Chatbot",
    page_icon="🤖",
    layout="wide"
)

st.title("🤖 Custom AI Document Chatbot")
st.caption("SkillsQuark – Task 3: Ask questions about your uploaded documents.")


def extract_pdf(file_bytes):
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name

    reader = PdfReader(tmp_path)
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(
                Document(
                    page_content=text,
                    metadata={"source": "uploaded PDF", "page": number}
                )
            )
    Path(tmp_path).unlink(missing_ok=True)
    return pages


def extract_txt(file_bytes, filename):
    text = file_bytes.decode("utf-8", errors="ignore")
    return [Document(page_content=text, metadata={"source": filename})]


def extract_csv(file_bytes, filename):
    with tempfile.NamedTemporaryFile(delete=False, suffix=".csv") as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name

    df = pd.read_csv(tmp_path)
    Path(tmp_path).unlink(missing_ok=True)

    text = df.to_csv(index=False)
    return [Document(page_content=text, metadata={"source": filename})]


def extract_docx(file_bytes, filename):
    try:
        from docx import Document as DocxDocument
    except ImportError:
        st.error("python-docx is not installed. Run: pip install -r requirements.txt")
        return []

    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name

    doc = DocxDocument(tmp_path)
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    Path(tmp_path).unlink(missing_ok=True)
    return [Document(page_content=text, metadata={"source": filename})]


def extract_documents(uploaded_files):
    documents = []

    for file in uploaded_files:
        raw = file.getvalue()
        suffix = Path(file.name).suffix.lower()

        if suffix == ".pdf":
            documents.extend(extract_pdf(raw))
        elif suffix == ".txt":
            documents.extend(extract_txt(raw, file.name))
        elif suffix == ".csv":
            documents.extend(extract_csv(raw, file.name))
        elif suffix == ".docx":
            documents.extend(extract_docx(raw, file.name))
        else:
            st.warning(f"Unsupported file skipped: {file.name}")

    return documents


def build_vector_store(documents):
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=150
    )
    chunks = splitter.split_documents(documents)

    embeddings = OpenAIEmbeddings(model="text-embedding-3-small")
    vector_store = FAISS.from_documents(chunks, embeddings)

    return vector_store, len(chunks)


def answer_question(vector_store, question):
    retriever = vector_store.as_retriever(search_kwargs={"k": 4})
    retrieved_docs = retriever.invoke(question)

    context_parts = []
    for doc in retrieved_docs:
        source = doc.metadata.get("source", "document")
        page = doc.metadata.get("page")
        label = f"{source}, page {page}" if page else source
        context_parts.append(f"[Source: {label}]\n{doc.page_content}")

    context = "\n\n".join(context_parts)

    prompt = ChatPromptTemplate.from_messages([
        (
            "system",
            """You are a document question-answering assistant.
Answer the user's question using only the supplied document context.
If the answer is not contained in the context, clearly say:
"I could not find that information in the uploaded documents."
Do not invent facts. Keep the answer clear and concise.
At the end, list the source names/pages used when available.

DOCUMENT CONTEXT:
{context}"""
        ),
        ("human", "{question}")
    ])

    model = ChatOpenAI(
        model="gpt-5-mini",
        temperature=0
    )

    chain = prompt | model
    response = chain.invoke({
        "context": context,
        "question": question
    })

    return response.content, retrieved_docs


def main():
    api_key = os.getenv("OPENAI_API_KEY")

    with st.sidebar:
        st.header("1. Upload documents")
        uploaded_files = st.file_uploader(
            "Upload PDF, TXT, CSV, or DOCX files",
            type=["pdf", "txt", "csv", "docx"],
            accept_multiple_files=True
        )

        if not api_key:
            st.warning(
                "OPENAI_API_KEY is not set. Create a .env file using .env.example."
            )

        process = st.button(
            "Build Knowledge Base",
            type="primary",
            disabled=not uploaded_files or not api_key
        )

    if "vector_store" not in st.session_state:
        st.session_state.vector_store = None

    if process:
        with st.spinner("Reading and indexing your documents..."):
            documents = extract_documents(uploaded_files)

            if not documents:
                st.error("No readable content was found.")
                return

            try:
                vector_store, count = build_vector_store(documents)
                st.session_state.vector_store = vector_store
                st.session_state.file_names = [f.name for f in uploaded_files]
                st.session_state.chunk_count = count
                st.success(f"Knowledge base created with {count} text chunks.")
            except Exception as exc:
                st.error(f"Could not build the knowledge base: {exc}")

    if st.session_state.vector_store is None:
        st.info(
            "Upload a document from the sidebar, click 'Build Knowledge Base', "
            "then ask a question below."
        )
        return

    st.subheader("2. Ask your documents")
    st.write(
        "Loaded: " + ", ".join(st.session_state.get("file_names", []))
    )

    question = st.text_input(
        "Your question",
        placeholder="What is the main objective of this document?"
    )

    if st.button("Ask AI"):
        if not question.strip():
            st.warning("Please enter a question.")
            return

        with st.spinner("Searching the documents and generating an answer..."):
            try:
                answer, sources = answer_question(
                    st.session_state.vector_store,
                    question
                )
                st.markdown("### Answer")
                st.write(answer)

                with st.expander("Retrieved document passages"):
                    for i, doc in enumerate(sources, start=1):
                        page = doc.metadata.get("page")
                        source = doc.metadata.get("source", "document")
                        label = f"{source} – page {page}" if page else source
                        st.markdown(f"**Passage {i}: {label}**")
                        st.write(doc.page_content)
            except Exception as exc:
                st.error(f"Question answering failed: {exc}")


if __name__ == "__main__":
    main()
